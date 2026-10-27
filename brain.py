"""brain.py — turn any uploaded file into Claude-ready content blocks.

Derived from HAL's hal_brain.process_uploads(), with two changes:
  * PDFs are sent natively (document blocks), so scanned/image PDFs and
    layouts are read by Claude itself and page numbers can be cited.
    Only PDFs over the API limits fall back to text with [PAGE n] markers.
  * Large images are downscaled so a phone photo cannot break the request.
"""
from __future__ import annotations

import base64
import io

_IMAGE_MIME = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "webp": "image/webp", "gif": "image/gif",
}
MAX_PDF_BYTES = 30 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_IMAGE_BYTES = 4 * 1024 * 1024
MAX_TEXT_CHARS = 120_000


def _ext(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def _read(item) -> tuple[str, bytes]:
    """Accept a Streamlit UploadedFile or a (name, bytes) tuple."""
    if isinstance(item, tuple):
        return item[0], item[1]
    name = getattr(item, "name", "file")
    try:
        item.seek(0)
    except Exception:
        pass
    return name, item.read()


def _pdf_pages(data: bytes) -> int | None:
    try:
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return len(doc)
    except Exception:
        return None


def _pdf_to_text(data: bytes) -> str:
    try:
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            parts = []
            for i, page in enumerate(doc, 1):
                t = page.get_text().strip()
                if t:
                    parts.append(f"[PAGE {i}]\n{t}")
        return "\n\n".join(parts)[:MAX_TEXT_CHARS] or "[PDF without extractable text]"
    except Exception as e:  # pragma: no cover - defensive
        return f"[Could not read PDF: {e}]"


def _docx_to_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)[:MAX_TEXT_CHARS] or "[Empty Word document]"


def _table_to_text(data: bytes, ext: str) -> str:
    import pandas as pd
    if ext in ("csv", "tsv"):
        df = pd.read_csv(io.BytesIO(data), sep="\t" if ext == "tsv" else ",")
        return df.to_csv(index=False)[:MAX_TEXT_CHARS]
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    out = [f"=== SHEET: {n} ===\n{df.to_csv(index=False)}" for n, df in sheets.items()]
    return "\n\n".join(out)[:MAX_TEXT_CHARS] or "[Empty spreadsheet]"


def _plain_to_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "cp1253", "latin-1"):
        try:
            return data.decode(enc)[:MAX_TEXT_CHARS]
        except Exception:
            continue
    return "[Could not decode text file]"


def _shrink_image(data: bytes) -> tuple[bytes, str]:
    """Return (bytes, media_type) small enough for the API."""
    from PIL import Image
    img = Image.open(io.BytesIO(data))
    img.thumbnail((2000, 2000))
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue(), "image/jpeg"


def _text_doc(name: str, text: str) -> dict:
    return {
        "type": "document",
        "title": name,
        "source": {"type": "text", "media_type": "text/plain", "data": text},
    }


def process_uploads(files) -> tuple[list, list[tuple[str, str]]]:
    """Return (blocks, summaries) where summaries = [(filename, icon)]."""
    blocks: list = []
    summaries: list[tuple[str, str]] = []

    for item in files or []:
        name, data = _read(item)
        ext = _ext(name)
        try:
            if ext in _IMAGE_MIME:
                media = _IMAGE_MIME[ext]
                if len(data) > MAX_IMAGE_BYTES:
                    data, media = _shrink_image(data)
                blocks.append({"type": "text", "text": f"[Image file: {name}]"})
                blocks.append({
                    "type": "image",
                    "source": {"type": "base64", "media_type": media,
                               "data": base64.standard_b64encode(data).decode()},
                })
                summaries.append((name, "🖼️"))
            elif ext == "pdf":
                pages = _pdf_pages(data)
                if len(data) <= MAX_PDF_BYTES and (pages is None or pages <= MAX_PDF_PAGES):
                    blocks.append({
                        "type": "document",
                        "title": name,
                        "source": {"type": "base64", "media_type": "application/pdf",
                                   "data": base64.standard_b64encode(data).decode()},
                    })
                    summaries.append((name, "📄"))
                else:
                    blocks.append(_text_doc(name, _pdf_to_text(data)))
                    summaries.append((name, "📄✂️"))
            elif ext == "docx":
                blocks.append(_text_doc(name, _docx_to_text(data)))
                summaries.append((name, "📝"))
            elif ext in ("xlsx", "xls", "csv", "tsv"):
                blocks.append(_text_doc(name, _table_to_text(data, ext)))
                summaries.append((name, "📊"))
            else:
                blocks.append(_text_doc(name, _plain_to_text(data)))
                summaries.append((name, "📃"))
        except Exception as e:
            summaries.append((name, f"⚠️ {e}"))
    return blocks, summaries
